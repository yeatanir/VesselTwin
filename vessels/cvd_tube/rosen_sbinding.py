"""DFT S-binding on the MoS2 Mo-zigzag edge from Rosen, Notestein & Snurr, J. Phys. Chem. C 122, 15318 (2018), DOI 10.1021/acs.jpcc.8b02524,
supporting-information workbook mos2_data.xlsx (path via env/arg; NOT committed, only the derived numbers in vessels/cvd_tube/rosen_mo_edge.json).
python -m vessels.cvd_tube.rosen_sbinding /path/to/mos2_data.xlsx  -> rewrites vessels/cvd_tube/rosen_mo_edge.json

What is computed (all theta_H = 0 column, i.e. no H on the edge):
 E_i   : sheet E_Mo_edge, column H-0, rows S-0 ... S-100 (xlsx rows 3-9, column B), eV, relative to the Mo-0/0 edge.
 F_i(T): E_i + F_vib(S*) from sheet S_vibrations_Mo_edge (harmonic, ZPE included; paper eqs 1-3, 6; same as the authors' mos2_data.py).
 dE_S  : differential S adsorption energy vs H2S/H2 (paper eq 19; Table 2): dE_S,i = (E_i - E_{i-1}) + E_H2 - E_H2S, E_H2/E_H2S from
         sheet Reference_species B2/B3 (-6.7755, -11.2255 eV). Reproduces Table 2 (checked: -2.91 -2.37 -0.93 -0.21 -0.08 -0.36).
 dF_S,i(T) = (F_i - F_{i-1}) + E_H2 - E_H2S (adds the S* vibrational free energy).
Conversion to our S2 vapour (ASSUMPTION, see dmu_S2_std): the edge sits in equilibrium with a gas of chemical potential mu_S = 1/2 mu_S2.
Rosen's Dmu_S* = mu_S* - (E_H2S - E_H2) (eq 15). Gas reaction H2S = H2 + 1/2 S2 with ideal-gas NIST/JANAF thermochemistry (no DFT) gives
 Dmu_S(S2 gas, T, p) = [mu0_H2S - mu0_H2](T) - (E_H2S - E_H2) + dG_r(T) + 1/2 kT ln(pS2 / 1 bar),
 [mu0 terms as the authors' get_mu_ref: H(T)-H(0K)+ZPE+E_DFT-TS, so the DFT energies cancel and only thermal parts remain].
Edge in equilibrium with S2 at p: dF_S,i(T) = Dmu_S  =>  p_eq,i = 1 bar exp(2 (dF_S,i - Dmu_S(T,1 bar)) / kT)."""
import json, os, sys
import numpy as np

KB = 8.6173303e-5  # eV/K, as in the authors' script
KJ2EV = 0.01036427
H_C = 6.626070040e-34 * 2.99792458e8 / 1.6021766208e-19 * 100  # eV per cm^-1
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f"{HERE}/rosen_mo_edge.json"
# NIST WebBook Shomate (A..H), copied by the authors into SI sheet Tabulated_data (rows 2-9 H2 <=1000 K, 12-19 H2 1000-2500 K, 32-39 H2S <=1400 K)
H2_LO = (33.066178, -11.363417, 11.432816, -2.772874, -0.158558, -9.980797, 172.707974, 0.0)
H2_HI = (18.563083, 12.257357, -2.859786, 0.268238, 1.97799, -1.147438, 156.288133, 0.0)
H2S_LO = (26.88412, 18.67809, 3.434203, -3.378702, 0.135882, -28.91211, 233.3747, -20.50202)
H2_H0K, H2S_H0K = -0.08775427409, -0.10324885774  # eV, sheet Tabulated_data rows 52, 55
H2_ZPE, H2S_ZPE = 0.2728, 0.3982372  # eV, rows 58, 61
THETA = [0, 1 / 6, 1 / 3, 1 / 2, 2 / 3, 5 / 6, 1]


def shomate(c, T):
    A, B, C, D, E, F, G, H = c
    t = T / 1000.0
    return A * t + B * t**2 / 2 + C * t**3 / 3 + D * t**4 / 4 - E / t + F - H, A * np.log(t) + B * t + C * t**2 / 2 + D * t**3 / 3 - E / (2 * t**2) + G


def mu0_thermal(sp, T):
    """H(T)-H(0K)+ZPE-T S, eV (authors' get_mu_ref without E_DFT)."""
    c, h0, zpe = (H2_LO if T <= 1000 else H2_HI, H2_H0K, H2_ZPE) if sp == "H2" else (H2S_LO, H2S_H0K, H2S_ZPE)
    h, s = shomate(c, T)
    return h * KJ2EV - h0 + zpe - T * s * KJ2EV / 1000.0


def dG_r(T):
    """H2S = H2 + 1/2 S2, 1 bar standard state, eV per H2S, from NIST/JANAF Shomate (S2: repo thermo, dfH 128.6003 kJ/mol)."""
    from vessels.cvd_tube import thermo as th
    G = lambda c, dHf: ((dHf + shomate(c, T)[0]) * 1e3 - T * shomate(c, T)[1]) / 96485.332
    T = float(T)
    g_h2 = G(H2_LO if T <= 1000 else H2_HI, 0.0)
    g_h2s = G(H2S_LO, -20.50202)
    g_s2 = G(th._S2, 128.6003)
    return g_h2 + 0.5 * g_s2 - g_h2s


def dmu_S2_std(T):
    """Dmu_S (Rosen eq 15 convention) of S2 gas at 1 bar. NOTE the DFT energies cancel: only thermal parts of H2S/H2 plus the experimental dG_r.
    mu0_H2S - mu0_H2 - (E_H2S - E_H2) = thermal part difference (ZPE, H-H0K, -TS)."""
    return mu0_thermal("H2S", T) - mu0_thermal("H2", T) + dG_r(T)


def f_vib(freqs_cm, T):
    e = np.array([f for f in freqs_cm if f and f == f]) * H_C
    zpe = 0.5 * e.sum()
    u = zpe + (e / (np.exp(e / (KB * T)) - 1)).sum()
    s = KB * (e / (KB * T * (np.exp(e / (KB * T)) - 1)) - np.log(1 - np.exp(-e / (KB * T)))).sum()
    return u - T * s


def build(xlsx, Ts=range(700, 1301, 50)):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    E = [wb["E_Mo_edge"].cell(r, 2).value for r in range(3, 10)]
    eh2, eh2s = wb["Reference_species"]["B2"].value, wb["Reference_species"]["B3"].value
    vib = [[c.value for c in row[1:] if c.value is not None] for row in wb["S_vibrations_Mo_edge"].iter_rows(min_row=1, max_row=6)]
    dE = [E[i] - E[i - 1] + eh2 - eh2s for i in range(1, 7)]
    dF = {}
    for T in Ts:
        F = [E[0]] + [E[i] + f_vib(vib[i - 1], T) for i in range(1, 7)]
        dF[str(T)] = [F[i] - F[i - 1] + eh2 - eh2s for i in range(1, 7)]
    return {"source": "Rosen, Notestein, Snurr, J. Phys. Chem. C 122, 15318 (2018), DOI 10.1021/acs.jpcc.8b02524; SI workbook mos2_data.xlsx",
            "E_Mo_edge_H0_eV": {"sheet": "E_Mo_edge", "column": "B (H-0)", "rows": "3-9 (S-0 .. S-100)", "values": E},
            "E_H2_eV": eh2, "E_H2S_eV": eh2s, "ref_cells": "Reference_species!B2, B3", "theta_S": THETA,
            "dE_S_eV_vs_H2S_H2": dE, "dE_S_matches_table2": [-2.91, -2.37, -0.93, -0.21, -0.08, -0.36],
            "dF_S_eV_vs_H2S_H2_by_T": dF, "vib_source": "sheet S_vibrations_Mo_edge rows 1-6 (S-17 .. S-100), harmonic, theta_H = 0"}


_D = None


def data():
    global _D
    if _D is None:
        _D = json.load(open(OUT))
    return _D


def dF_S(T):  # ponytail: clamps to the stored 700-1300 K grid outside it
    """Differential S free energies (6 steps, eV, vs H2S/H2) at T, linear interpolation in T over the stored grid."""
    d = data()["dF_S_eV_vs_H2S_H2_by_T"]
    Ts = sorted(float(k) for k in d)
    arr = np.array([d[str(int(t))] for t in Ts])
    return np.array([np.interp(T, Ts, arr[:, i]) for i in range(arr.shape[1])])


def p_eq_steps(T, steps=slice(3, 6)):
    """Equilibrium S2 pressure (Pa) of each S-coverage step: 1 bar exp(2 (dF_S - Dmu_S2_std)/kT)."""
    return 1e5 * np.exp(np.clip(2 * (dF_S(T)[steps] - dmu_S2_std(T)) / (KB * T), -700, 700))


if __name__ == "__main__":
    j = build(sys.argv[1])
    json.dump(j, open(OUT, "w"), indent=1)
    print(np.round(j["dE_S_eV_vs_H2S_H2"], 2), "Table 2 printed:", j["dE_S_matches_table2"])
